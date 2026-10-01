import io
from datetime import UTC, datetime

from openpyxl import load_workbook

from app.exporters import ExportColumn, export_bytes

COLUMNS = [ExportColumn("name", "Mesa"), ExportColumn("situation", "Situação"),
           ExportColumn("late_days", "Dias de atraso"), ExportColumn("created_at", "Aberto em")]
ROWS = [
    {"name": "SUPORTE", "situation": "open", "late_days": 10.5,
     "created_at": datetime(2026, 9, 29, 15, 0, tzinfo=UTC)},
    {"name": "INFRA", "situation": "canceled", "late_days": None, "created_at": None},
]
TZ = "America/Sao_Paulo"


def test_csv_is_excel_friendly() -> None:
    text = export_bytes("csv", ROWS, COLUMNS, "t", TZ).decode("utf-8")
    lines = text.lstrip("﻿").splitlines()
    assert text.startswith("﻿")
    assert lines[0] == "Mesa;Situação;Dias de atraso;Aberto em"
    assert lines[1] == "SUPORTE;Aberto;10,5;29/09/2026 12:00"
    assert lines[2] == "INFRA;Cancelado;;"


def test_xlsx_has_header_and_local_naive_dates() -> None:
    content = export_bytes("xlsx", ROWS, COLUMNS, "Um título bem longo que passa de 31 chars", TZ)
    sheet = load_workbook(io.BytesIO(content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Mesa", "Situação", "Dias de atraso", "Aberto em")
    assert rows[1] == ("SUPORTE", "Aberto", 10.5, datetime(2026, 9, 29, 12, 0))  # noqa: DTZ001 - Excel cells are naive local time
    assert len(rows) == 3
    assert len(sheet.title) <= 31


def test_description_is_exported_as_plain_text() -> None:
    columns = [ExportColumn("ticket_number", "Nº"), ExportColumn("description", "Descrição")]
    rows = [{"ticket_number": 1, "description": "<p>Bom dia</p><p>Segue &amp; anexo</p>"},
            {"ticket_number": 2, "description": None}]
    sheet = load_workbook(io.BytesIO(export_bytes("xlsx", rows, columns, "t", TZ))).active
    assert list(sheet.iter_rows(values_only=True))[1:] == [(1, "Bom dia\nSegue & anexo"), (2, None)]


def test_xlsx_cuts_text_above_the_excel_cell_limit() -> None:
    columns = [ExportColumn("description", "Descrição")]
    content = export_bytes("xlsx", [{"description": "a" * 40_000}], columns, "t", TZ)
    cell = list(load_workbook(io.BytesIO(content)).active.iter_rows(values_only=True))[1][0]
    assert len(cell) == 32_767 and cell.endswith("…(cortado)")
