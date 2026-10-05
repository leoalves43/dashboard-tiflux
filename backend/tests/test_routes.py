"""HTTP-level checks: filters must parse alongside endpoint-specific query params (regression)."""

import io
from collections.abc import Iterator

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine

from app.config import Settings
from app.main import create_app
from app.routes.deps import get_context, get_tiflux_source
from app.sla import QueryContext
from tests.conftest import NOW, TEST_SCHEMA
from tests.fakes import FakeActivitySource


@pytest.fixture(scope="module")
def client(engine: Engine) -> Iterator[TestClient]:
    app = create_app(Settings(schema_name=TEST_SCHEMA))  # type: ignore[call-arg]
    # Fixed clock like the query tests: late/aging buckets used the real now and drifted
    # (test_bucket_export broke on 2026-10-05 when ticket A crossed into 16–30 days).
    app.dependency_overrides[get_context] = lambda: QueryContext(now=NOW)
    with TestClient(app) as test_client:
        yield test_client


def test_timeseries_accepts_filters_and_granularity(client: TestClient) -> None:
    response = client.get("/api/timeseries", params={"granularity": "day", "desk_ids": [11]})
    assert response.status_code == 200, response.text
    assert response.json() == [{"period": "2026-09-25", "created": 1, "solved": 0}]


def test_tickets_accepts_filters_sort_and_paging(client: TestClient) -> None:
    params = {"situations": ["closed"], "sort": "ticket_number", "direction": "asc", "page_size": 1}
    body = client.get("/api/tickets", params=params).json()
    assert body["total"] == 2
    assert [row["ticket_number"] for row in body["rows"]] == [3]


def test_invalid_sort_is_422_with_value(client: TestClient) -> None:
    response = client.get("/api/tickets", params={"sort": "drop table"})
    assert response.status_code == 422 and "drop table" in response.text


def test_ticket_export_matches_filtered_rows(client: TestClient) -> None:
    response = client.get("/api/export/tickets/xlsx", params={"situations": ["open"], "sort": "late_days"})
    assert response.status_code == 200
    rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
    assert len(rows) == 1 + 3  # header + open tickets A, B, F



def test_ticket_export_keeps_chosen_columns_and_order(client: TestClient) -> None:
    # Sorting by a column left out of the export must still work (spec 004).
    params = {"sort": "late_days", "columns": ["title", "ticket_number"]}
    response = client.get("/api/export/tickets/xlsx", params=params)
    assert response.status_code == 200, response.text
    header = next(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
    assert header == ("Título", "Nº")


def test_ticket_export_rejects_unknown_column(client: TestClient) -> None:
    response = client.get("/api/export/tickets/csv", params={"columns": ["bogus"]})
    assert response.status_code == 422 and "bogus" in response.text

def test_breakdown_export_csv(client: TestClient) -> None:
    response = client.get("/api/export/breakdown/desk/csv", params={"sla": ["late"]})
    lines = response.content.decode("utf-8-sig").splitlines()
    assert lines[0].startswith("Mesa;Total;")
    assert lines[1].startswith("SUPORTE;2;")


def test_timeseries_export(client: TestClient) -> None:
    response = client.get("/api/export/timeseries/csv", params={"granularity": "month"})
    lines = response.content.decode("utf-8-sig").splitlines()
    assert lines == ["Período (início);Abertos;Resolvidos", "2026-08-01;1;0", "2026-09-01;5;2"]


def test_bucket_export(client: TestClient) -> None:
    response = client.get("/api/export/buckets/late/xlsx")
    rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
    assert ("8–15 dias", 1) in rows and len(rows) == 6
    assert client.get("/api/export/buckets/nope/csv").status_code == 404


def test_ticket_summary_route_returns_404_for_unknown_ticket(client: TestClient) -> None:
    assert client.get("/api/tickets/999999").status_code == 404
    assert client.get("/api/tickets/1").json()["ticket_number"] == 1


@pytest.mark.parametrize(("source", "status"), [
    (FakeActivitySource(lists={}), 200),
    (FakeActivitySource(lists={}, missing=True), 404),
    (FakeActivitySource(lists={}, failing=True), 502),
])
def test_ticket_activity_route_maps_tiflux_outcomes(client: TestClient, source: FakeActivitySource,
                                                    status: int) -> None:
    client.app.dependency_overrides[get_tiflux_source] = lambda: source  # type: ignore[attr-defined]
    try:
        responses = [client.get(f"/api/tickets/7/{part}") for part in ("activity", "description")]
    finally:
        client.app.dependency_overrides.pop(get_tiflux_source)  # type: ignore[attr-defined]
    # description of a ticket with no detail payload is 404 in the fake, so only check failure statuses there
    assert responses[0].status_code == status, responses[0].text
    assert responses[1].status_code == (404 if status == 200 else status), responses[1].text
